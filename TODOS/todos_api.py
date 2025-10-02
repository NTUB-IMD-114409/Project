# blueprints/todos_api.py
# -*- coding: utf-8 -*-
"""
API 一次到位（待辦專用）：
- POST /api/generate_todos        只整理前端勾選的逐字稿，抽取待辦事項
- GET  /api/todos/<meeting_id>    讀取 latest.json
- GET  /api/download_todos/<id>   下載 Word 檔（依 latest.json 產生）

需求套件：
  pip install pdfminer.six python-docx
"""

from __future__ import annotations

import os
import io
import re
import json
from datetime import datetime
from typing import List, Dict, Any, Optional, Tuple

from flask import Blueprint, request, jsonify, send_file

# 你的 LLM 包裝
from .services.llm_client import run_llm, LLM_CTX, MAX_OUT

todos_api = Blueprint("todos_api", __name__)

# ----------------------------- 路徑與 I/O -----------------------------

BASE_DIR = os.getcwd()
UPLOAD_ROOT = os.path.join(BASE_DIR, "uploads")
TODOS_ROOT = os.path.join(UPLOAD_ROOT, "todos")

def _ensure_dir(p: str) -> None:
    os.makedirs(p, exist_ok=True)

def _safe_abs_upload_path(rel_path: str) -> str:
    rel = rel_path.lstrip("/").replace("..", "")
    abs_path = os.path.normpath(os.path.join(UPLOAD_ROOT, rel))
    if not abs_path.startswith(os.path.abspath(UPLOAD_ROOT)):
        raise ValueError("非法路徑")
    return abs_path

def _todos_dir(meeting_id: int) -> str:
    return os.path.join(TODOS_ROOT, str(meeting_id))

def save_todos_json(meeting_id: int, data: Dict[str, Any]) -> None:
    """同時存 timestamp 檔與 latest.json"""
    d = _todos_dir(meeting_id)
    _ensure_dir(d)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    with open(os.path.join(d, f"todos_{ts}.json"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    with open(os.path.join(d, "latest.json"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def load_latest_todos(meeting_id: int) -> Optional[Dict[str, Any]]:
    p = os.path.join(_todos_dir(meeting_id), "latest.json")
    if not os.path.isfile(p):
        return None
    try:
        with open(p, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        print("[todos_api] load_latest_todos error:", e)
        return None

# ----------------------------- 逐字稿抽文字 -----------------------------

def extract_text_from_file(abs_path: str) -> str:
    p = abs_path.lower()
    try:
        if p.endswith((".txt", ".md", ".log")):
            with open(abs_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        if p.endswith(".docx"):
            from docx import Document
            doc = Document(abs_path)
            return "\n".join([para.text for para in doc.paragraphs])

        if p.endswith(".pdf"):
            try:
                from pdfminer.high_level import extract_text
                return extract_text(abs_path) or ""
            except Exception as e:
                print("[todos_api] pdfminer failed, try pypdf:", e)
                try:
                    from pypdf import PdfReader
                    texts = []
                    reader = PdfReader(abs_path)
                    for page in reader.pages:
                        texts.append(page.extract_text() or "")
                    return "\n".join(texts)
                except Exception as e2:
                    print("[todos_api] pypdf failed:", e2)
                    return ""

        if p.endswith((".srt", ".vtt")):
            with open(abs_path, "r", encoding="utf-8", errors="ignore") as f:
                raw = f.read()
            raw = re.sub(r"\d{2}:\d{2}:\d{2}[,\.]\d{3}\s*-->\s*\d{2}:\d{2}:\d{2}[,\.]\d{3}", "", raw)
            raw = re.sub(r"^\s*\d+\s*$", "", raw, flags=re.MULTILINE)
            return raw

    except Exception as e:
        print("[todos_api] extract_text_from_file error:", abs_path, e)
    return ""

def load_selected_transcripts_text(meeting_id: int, rel_paths: List[str]) -> str:
    pieces: List[str] = []
    for rel in rel_paths:
        try:
            abs_path = _safe_abs_upload_path(rel)
            if not os.path.isfile(abs_path):
                print("[todos_api] missing file:", abs_path)
                continue
            txt = extract_text_from_file(abs_path)
            if txt and txt.strip():
                pieces.append(txt)
            else:
                print("[todos_api] empty text parsed:", abs_path)
        except Exception as e:
            print("[todos_api] load_selected_transcripts_text error:", rel, e)
    return "\n\n".join(pieces).strip()

# ----------------------------- 分塊與合併 -----------------------------

def approx_tokens(s: str) -> int:
    return max(1, len(s) // 4)

def split_for_ctx(text: str, budget_tokens: int) -> List[str]:
    if approx_tokens(text) <= budget_tokens:
        return [text]
    out, cur, cur_tokens = [], [], 0
    lines = text.splitlines()
    for ln in lines:
        t = approx_tokens(ln) + 1
        if cur_tokens + t > budget_tokens and cur:
            out.append("\n".join(cur).strip())
            cur, cur_tokens = [], 0
        cur.append(ln)
        cur_tokens += t
    if cur:
        out.append("\n".join(cur).strip())
    if not out:
        chunks = []
        step = max(1, budget_tokens * 4)
        for i in range(0, len(text), step):
            chunks.append(text[i:i + step])
        return chunks
    return out

def merge_partials(items: List[Dict[str, Any]]) -> Dict[str, Any]:
    todos = []
    for it in items:
        if not it:
            continue
        todos += it.get("todos", []) or []
    return {"todos": todos}

# ----------------------------- Prompt 與清洗 -----------------------------

BASE_PROMPT = """你是專業會議紀錄助理。請依【逐字稿】找出所有待辦事項，只輸出「單一 JSON 物件」。不得輸出任何多餘文字或標記。

JSON 結構：
{
  "todos": [
    {"content":"string"}
  ]
}

規則：
- 只要明確的待辦事項，逐字稿中提到的即可。
- 不要輸出負責人、期限。
- 所有輸出必須為 JSON；不要 ```、不要解釋。

【逐字稿】
{TEXT}
"""

STRICT_REPROMPT = """嚴格只輸出有效 JSON（單一物件），不可包含任何額外文字或標記。
結構必須為：
{"todos":[{"content":"string"}]}

【逐字稿】
{TEXT}
"""

def _strip_markdown_and_extract_json(s: str) -> str:
    s = re.sub(r"```json\s*([\s\S]*?)```", r"\1", s, flags=re.IGNORECASE)
    s = re.sub(r"```([\s\S]*?)```", r"\1", s)
    m = re.search(r"\{[\s\S]*\}", s)
    return (m.group(0) if m else s).strip()

def _safe_json_loads(s: str) -> Optional[Dict[str, Any]]:
    try:
        return json.loads(s)
    except Exception:
        m = re.search(r"\{[\s\S]*\}", s)
        if m:
            try:
                return json.loads(m.group(0))
            except Exception:
                return None
        return None

def _llm_todos_once(text: str, strict: bool = False) -> Tuple[Optional[Dict[str, Any]], str]:
    prompt = (STRICT_REPROMPT if strict else BASE_PROMPT).replace("{TEXT}", text)
    out = run_llm(prompt)
    cleaned = _strip_markdown_and_extract_json(out)
    obj = _safe_json_loads(cleaned)
    return obj, cleaned

# ----------------------------- Routes -----------------------------

@todos_api.route("/api/generate_todos", methods=["POST"])
def generate_todos():
    data = request.get_json(force=True, silent=True) or {}
    meeting_id = int(data.get("meeting_id") or 0)
    selected: List[str] = data.get("transcripts", [])

    if not meeting_id:
        return jsonify({"success": False, "message": "缺少 meeting_id"}), 400
    if not selected:
        return jsonify({"success": False, "message": "請勾選至少一個逐字稿檔案"}), 400

    full_text = load_selected_transcripts_text(meeting_id, selected)
    if not full_text:
        return jsonify({"success": False, "message": "勾選的檔案無法讀取文字"}), 400

    input_budget = max(512, LLM_CTX - (MAX_OUT + 512))
    chunks = [full_text] if approx_tokens(full_text) <= input_budget else split_for_ctx(full_text, input_budget)

    partials: List[Dict[str, Any]] = []
    cleaned_samples: List[str] = []

    for ch in chunks:
        try:
            obj, cleaned = _llm_todos_once(ch, strict=False)
            if not obj:
                obj, cleaned = _llm_todos_once(ch, strict=True)
            cleaned_samples.append(cleaned)
            partials.append(obj or {"todos": []})
        except Exception as e:
            print("[todos_api] LLM error:", e)
            partials.append({"todos": []})

    result = merge_partials(partials)

    try:
        save_todos_json(meeting_id, result)
    except Exception as e:
        print("[todos_api] save_todos_json error:", e)

    return jsonify({"success": True, "data": result})

@todos_api.route("/api/todos/<int:meeting_id>", methods=["GET"])
def get_latest_todos(meeting_id: int):
    data = load_latest_todos(meeting_id)
    return jsonify({"success": True, "data": data})

@todos_api.route("/api/download_todos/<int:meeting_id>", methods=["GET"])
def download_todos(meeting_id: int):
    data = load_latest_todos(meeting_id)
    if not data:
        return jsonify({"success": False, "message": "沒有可下載的待辦事項"}), 404

    from docx import Document
    from docx.shared import Pt

    doc = Document()
    h1 = doc.add_heading("會議待辦事項", level=1)
    h1.style.font.size = Pt(18)

    # 待辦
    doc.add_heading("📝 待辦事項", level=2)
    items = data.get("todos") or []
    if not items:
        doc.add_paragraph("(無)")
    else:
        for it in items:
            content = it.get("content") or "未命名"
            doc.add_paragraph(f"- {content}")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    fname = f"meeting_{meeting_id}_todos.docx"
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=fname,
    )