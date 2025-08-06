# speech_transcribe.py

from flask import Blueprint, request, jsonify, send_from_directory
import os
from datetime import datetime
from werkzeug.utils import secure_filename
from opencc import OpenCC
from transformers import AutoTokenizer, AutoModelForSeq2SeqLM
from docx import Document
from .whisper_core import transcribe_audio_with_cli, extract_srt_lines_with_timestamp
from .cleaner import clean_text
from .punctuate import basic_punctuation, bart_punctuation

# === 模型初始化 ===
tokenizer = AutoTokenizer.from_pretrained("fnlp/bart-base-chinese")
bart_model = AutoModelForSeq2SeqLM.from_pretrained("fnlp/bart-base-chinese")

# === 資料夾初始化 ===
UPLOAD_FOLDER = os.path.join("uploads", "audio")  # ✔ uploads/audio/
TRANSCRIPT_FOLDER = os.path.join("uploads", "transcripts")  # ✔ uploads/transcripts/

UPLOAD_FOLDER = os.path.join("uploads", "audio")
TRANSCRIPT_FOLDER = os.path.join("uploads", "transcripts")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(TRANSCRIPT_FOLDER, exist_ok=True)

speech_bp = Blueprint('speech', __name__)

# === 上傳音訊檔 ===
@speech_bp.route("/upload_audio", methods=["POST"])
def upload_audio():
    file = request.files.get("file")
    if not file:
        return jsonify({"error": "未提供檔案"}), 400

    filename = secure_filename(file.filename)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    saved_filename = f"{timestamp}_{filename}"

    save_path = os.path.join(UPLOAD_FOLDER, saved_filename)
    file.save(save_path)

    return jsonify({"filename": saved_filename}), 200

# === Whisper 轉文字 + 清洗 + 補標點 ===
@speech_bp.route("/whisper", methods=["POST"])
def whisper_transcribe():
    try:
        data = request.get_json()
        filename = data.get("filename")
        use_basic = data.get("use_basic", False)

        if not filename:
            return jsonify({"error": "未提供檔名"}), 400

        file_path = os.path.join(UPLOAD_FOLDER, filename)
        if not os.path.exists(file_path):
            return jsonify({"error": "檔案不存在"}), 404

        # 1️⃣ 執行 Whisper CLI（產出 txt/srt/vtt/tsv/json）
        transcribe_audio_with_cli(file_path, model_size="medium", lang="zh")

        # 2️⃣ 讀取 Whisper CLI 輸出的 .txt
        txt_raw_path = os.path.splitext(file_path)[0] + ".txt"
        if not os.path.exists(txt_raw_path):
            raise FileNotFoundError(f"找不到 Whisper CLI 輸出的 txt 檔案：{txt_raw_path}")
        with open(txt_raw_path, "r", encoding="utf-8") as f:
            raw_text = f.read().strip()

        # 3️⃣ 清洗雜訊與贅詞
        cleaned_text = clean_text(raw_text)

        # 4️⃣ 簡轉繁
        converter = OpenCC("s2t")
        traditional_text = converter.convert(cleaned_text)

        # 5️⃣ 補標點
        if use_basic:
            polished_text = basic_punctuation(traditional_text)
        else:
            polished_text = bart_punctuation(traditional_text)

        # 6️⃣ 儲存為 txt
        base_name = os.path.splitext(filename)[0]
        txt_filename = f"{base_name}_polished.txt"
        txt_path = os.path.join(TRANSCRIPT_FOLDER, txt_filename)
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(polished_text)

        # 7️⃣ 儲存為 docx
        docx_filename = f"{base_name}_polished.docx"
        docx_path = os.path.join(TRANSCRIPT_FOLDER, docx_filename)
        doc = Document()
        doc.add_heading("逐字稿紀錄", level=1)
        doc.add_paragraph(polished_text)
        doc.save(docx_path)

        # 8️⃣ 顯示逐句內容（從 .srt 取時間＋句子）
        lines = extract_srt_lines_with_timestamp(file_path)
        print("=== 逐句轉寫內容（含時間） ===")
        for idx, line in enumerate(lines, 1):
            print(f"{idx}. {line}")
        print("=== End ===")

        return jsonify({
            "result": polished_text,
            "lines": lines,
            "download_txt": f"/download/{txt_filename}",
            "download_docx": f"/download/{docx_filename}"
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

# === 檔案下載 ===
@speech_bp.route("/download/<filename>")
def download_file(filename):
    return send_from_directory(TRANSCRIPT_FOLDER, filename, as_attachment=True)
