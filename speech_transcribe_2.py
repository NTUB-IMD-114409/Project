# blueprints/speech_transcribe.py
# -*- coding: utf-8 -*-

from flask import Blueprint, request, jsonify, send_from_directory
import os
from datetime import datetime
from werkzeug.utils import secure_filename
from opencc import OpenCC
from docx import Document

from .whisper_core import transcribe_audio_with_cli, extract_srt_lines_with_timestamp
from .cleaner import clean_text
from .punctuate import basic_punctuation, bart_punctuation

# === 路徑設定 ===
UPLOAD_FOLDER = os.path.join("uploads", "audio")
TRANSCRIPT_FOLDER = os.path.join("uploads", "transcripts")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(TRANSCRIPT_FOLDER, exist_ok=True)

# === 轉換器（簡轉繁）===
_converter = OpenCC("s2t")

speech_bp = Blueprint("speech", __name__)


# ========== 工具 ==========
def _merge_lines_text(lines):
    """
    從 extract_srt_lines_with_timestamp 的結果中，移除時間與分隔符，合併為純文字。
    例：'00:00:00,000 --> 00:00:05,000 | 大家好' -> '大家好'
    """
    if not lines:
        return ""
    out = []
    for ln in lines:
        # 拆掉 '時間 | 內容'
        parts = ln.split("|", 1)
        if len(parts) == 2:
            out.append(parts[1].strip())
        else:
            out.append(ln.strip())
    return "\n".join(out).strip()


# ========== 上傳音訊 ==========
@speech_bp.route("/upload_audio", methods=["POST"])
def upload_audio():
    file = request.files.get("file")
    if not file:
        return jsonify({"error": "未提供檔案"}), 400

    filename = secure_filename(file.filename)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    saved_filename = f"{timestamp}_{filename}"

    save_path = os.path.join(UPLOAD_FOLDER, saved_filename)
    file.save(save_path)

    return jsonify({"filename": saved_filename}), 200


# ========== Whisper 轉文字 + 清洗 + 補標點 ==========
@speech_bp.route("/whisper", methods=["POST"])
def whisper_transcribe():
    try:
        data = request.get_json(silent=True) or {}
        filename = data.get("filename")
        # 前端可傳 use_basic=True 直接走規則式標點，避免 BART 造成短輸出
        use_basic = bool(data.get("use_basic", False))
        # meeting_id 可傳但此處不強制使用
        # meeting_id = data.get("meeting_id")

        if not filename:
            return jsonify({"error": "未提供檔名"}), 400

        file_path = os.path.join(UPLOAD_FOLDER, filename)
        if not os.path.exists(file_path):
            return jsonify({"error": "檔案不存在"}), 404

        # 1) Whisper CLI 轉錄（輸出 txt/srt/vtt/tsv/json）
        transcribe_audio_with_cli(file_path, model_size="medium", lang="zh")

        # 2) 讀取 Whisper 輸出的 .txt（原始粗文本）
        txt_raw_path = os.path.splitext(file_path)[0] + ".txt"
        if not os.path.exists(txt_raw_path):
            raise FileNotFoundError(f"找不到 Whisper CLI 輸出的 txt 檔案：{txt_raw_path}")

        with open(txt_raw_path, "r", encoding="utf-8") as f:
            raw_text = f.read().strip()

        # 3) 清洗雜訊與贅詞
        cleaned_text = clean_text(raw_text)

        # 4) 簡轉繁
        traditional_text = _converter.convert(cleaned_text)

        # 5) 補標點：basic 或 BART（含自動回退保險）
        if use_basic:
            polished_text = basic_punctuation(traditional_text)
        else:
            try:
                polished_text = bart_punctuation(traditional_text)
            except Exception as _:
                # 若 BART 失敗（顯存不足、模型未載入等），回退 basic
                polished_text = basic_punctuation(traditional_text)

        # 若結果異常短（可能被模型截短），再保險一次回退 basic
        if len(polished_text.strip()) < 50 and len(traditional_text.strip()) >= 50:
            polished_text = basic_punctuation(traditional_text)

        # 6) 準備逐句（SRT）資料，前端可用來備援
        lines = extract_srt_lines_with_timestamp(file_path)

        # 如果結果仍然非常短，嘗試用 SRT lines 合併作為備援
        if len(polished_text.strip()) < 20 and lines:
            merged = _merge_lines_text(lines)
            if len(merged) > len(polished_text):
                polished_text = merged

        # 7) 輸出 txt 檔
        base_name = os.path.splitext(filename)[0]
        txt_filename = f"{base_name}_polished.txt"
        txt_path = os.path.join(TRANSCRIPT_FOLDER, txt_filename)
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(polished_text)

        # 8) 輸出 docx 檔
        docx_filename = f"{base_name}_polished.docx"
        docx_path = os.path.join(TRANSCRIPT_FOLDER, docx_filename)
        doc = Document()
        doc.add_heading("逐字稿紀錄", level=1)
        for para in polished_text.splitlines():
            doc.add_paragraph(para)
        doc.save(docx_path)

        # 9) 回傳（加上長度資訊，方便前端/你排錯）
        return jsonify({
            "result": polished_text,
            "lines": lines,
            "download_txt": f"/download/{txt_filename}",
            "download_docx": f"/download/{docx_filename}",
            "len_result": len(polished_text),
            "len_raw": len(traditional_text),
        }), 200

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


# ========== 檔案下載 ==========
@speech_bp.route("/download/<filename>")
def download_file(filename):
    return send_from_directory(TRANSCRIPT_FOLDER, filename, as_attachment=True)
