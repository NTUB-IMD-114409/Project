# blueprints/speech_transcribe.py
# -*- coding: utf-8 -*-

from __future__ import annotations
from flask import Blueprint, request, jsonify, send_from_directory, session, Response, stream_with_context
import os
import json
import time
import re
import subprocess
from datetime import datetime
from werkzeug.utils import secure_filename
from opencc import OpenCC
from docx import Document

# --- 專案內部匯入 ---
from db import get_db  # ✅ 寫入 DB
from .whisper_core import (
    transcribe_audio_with_cli,
    extract_srt_lines_with_timestamp,
    _resolve_whisper_cmd,          # ✅ 用你現有的偵測
    _detect_device_for_cli,        # ✅ 自動選 cuda/cpu
    _cpu_threads,                  # ✅ 與 CLI 旗標一致
)
from .cleaner import clean_text
from .punctuate import basic_punctuation, bart_punctuation

# 逐行 SRT 標點（保留時間碼）
from .openai_tools.openai_srt_punctuator import punctuate_srt_lines
# SRT 工具（單行/標準 SRT 轉換）
from .openai_tools.srt_utils import parse_lines, write_entries, list_to_srt_blocks

# 整段文本用 GPT 標點
from .openai_tools.gpt_punctuate import punctuate_with_gpt

# === 路徑設定 ===
UPLOAD_FOLDER = os.path.join("uploads", "audio")  # 音訊暫存
MEETING_ROOT = "uploads"                           # 會議資料夾根目錄

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(MEETING_ROOT, exist_ok=True)

# === 簡轉繁 ===
_converter = OpenCC("s2t")

speech_bp = Blueprint("speech", __name__)

# ========== 工具 ==========
def _merge_lines_text(lines: list[str]) -> str:
    """把 SRT 單行『時間 --> 時間 | 內容』合併成純文字。"""
    if not lines:
        return ""
    out = []
    for ln in lines:
        parts = ln.split("|", 1)
        out.append(parts[1].strip() if len(parts) == 2 else ln.strip())
    return "\n".join(out).strip()

def _meeting_dir(meeting_id: int) -> str:
    """回傳 uploads/meeting_<id>，若無則建立。"""
    path = os.path.join(MEETING_ROOT, f"meeting_{int(meeting_id)}")
    os.makedirs(path, exist_ok=True)
    return path

def _upsert_file_record(meeting_id: int, file_name: str, rel_path: str, file_type: str, uploaded_by: int):
    """
    以 (meeting_id, file_name) 當自然鍵做 upsert。
    若 DB 的 files 有 UNIQUE KEY (meeting_id, file_name)，可調整成 ON DUPLICATE KEY UPDATE。
    """
    conn = get_db()
    cur = conn.cursor(dictionary=True)
    try:
        cur.execute(
            "SELECT id FROM files WHERE meeting_id=%s AND file_name=%s LIMIT 1",
            (meeting_id, file_name)
        )
        row = cur.fetchone()
        if row:
            cur.execute(
                """
                UPDATE files
                SET file_path=%s, file_type=%s, uploaded_at=NOW(), uploaded_by=%s
                WHERE id=%s
                """,
                (rel_path, file_type, uploaded_by, row["id"])
            )
        else:
            cur.execute(
                """
                INSERT INTO files (meeting_id, file_name, file_path, file_type, uploaded_at, uploaded_by)
                VALUES (%s, %s, %s, %s, NOW(), %s)
                """,
                (meeting_id, file_name, rel_path, file_type, uploaded_by)
            )
        conn.commit()
    finally:
        cur.close()
        conn.close()

# ========== 上傳音訊 ==========
@speech_bp.route("/upload_audio", methods=["POST"])
def upload_audio():
    file = request.files.get("file")
    if not file:
        return jsonify({"success": False, "error": "未提供檔案"}), 400

    filename = secure_filename(file.filename)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    saved_filename = f"{timestamp}_{filename}"

    save_path = os.path.join(UPLOAD_FOLDER, saved_filename)
    file.save(save_path)

    return jsonify({"success": True, "filename": saved_filename}), 200

# ========== Whisper 串流（SSE）：邊跑邊吐字幕行 ==========
@speech_bp.route("/whisper/stream")
def whisper_stream():
    """
    前端以 EventSource 連線：
      GET /whisper/stream?filename=xxx&meeting_id=123

    事件型別：
      - segment: { ts, text }            # 每出一段就送一次
      - meta:    { download_txt: ... }   # 完成後提供下載連結或錯誤
      - done:    {}                      # 串流結束
      - log:     { line: "..."}          # （可選）CLI 日誌
    """
    filename = (request.args.get("filename") or "").strip()
    meeting_id = (request.args.get("meeting_id") or "").strip()

    if not filename or not meeting_id:
        # 用標準 200 + SSE error 事件回去，避免 EventSource 直接報錯中斷
        def bad():
            yield f"event: meta\ndata: {json.dumps({'error':'缺少參數 filename 或 meeting_id'})}\n\n"
            yield "event: done\ndata: {}\n\n"
        return Response(stream_with_context(bad()), headers={
            "Content-Type": "text/event-stream; charset=utf-8",
            "Cache-Control": "no-cache, no-transform",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        })

    try:
        meeting_id = int(meeting_id)
    except Exception:
        def bad2():
            yield f"event: meta\ndata: {json.dumps({'error':'meeting_id 必須為整數'})}\n\n"
            yield "event: done\ndata: {}\n\n"
        return Response(stream_with_context(bad2()), headers={
            "Content-Type": "text/event-stream; charset=utf-8",
            "Cache-Control": "no-cache, no-transform",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        })

    file_path = os.path.join(UPLOAD_FOLDER, filename)
    if not os.path.exists(file_path):
        def bad3():
            yield f"event: meta\ndata: {json.dumps({'error':'檔案不存在'})}\n\n"
            yield "event: done\ndata: {}\n\n"
        return Response(stream_with_context(bad3()), headers={
            "Content-Type": "text/event-stream; charset=utf-8",
            "Cache-Control": "no-cache, no-transform",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        })

    base, _ = os.path.splitext(file_path)
    srt_path = base + ".srt"
    txt_raw_path = base + ".txt"

    user_info = session.get("user") or {}
    try:
        uploader_id = int(user_info.get("id", 0) or 0)
    except Exception:
        uploader_id = 0

    def gen():
        # 心跳：避免 Proxy/瀏覽器關連線
        yield ": heartbeat\n\n"

        # 先檢查 ffmpeg 是否存在
        try:
            subprocess.run(["ffmpeg", "-version"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        except Exception:
            yield f"event: meta\ndata: {json.dumps({'error':'ffmpeg 不可用'})}\n\n"
            yield "event: done\ndata: {}\n\n"
            return

        # 解析 whisper 指令與裝置
        try:
            whisper_cmd = _resolve_whisper_cmd()
        except Exception as e:
            yield f"event: meta\ndata: {json.dumps({'error': f'找不到 whisper CLI：{str(e)}'})}\n\n"
            yield "event: done\ndata: {}\n\n"
            return

        device = _detect_device_for_cli(min_free_gb=4.0, prefer_gpu=True)
        threads = str(_cpu_threads())

        # 組 CLI 旗標（與你原來一致，但 verbose=True 以便 stdout 逐行印出）
        cmd = [
            whisper_cmd,
            file_path,
            "--model", "medium",
            "--language", "zh",
            "--task", "transcribe",
            "--output_format", "all",
            "--output_dir", os.path.dirname(file_path),
            "--verbose", "True",
            "--device", device,
            "--threads", threads,
            "--condition_on_previous_text", "False",
        ]
        if device == "cpu":
            cmd += ["--fp16", "False"]

        # 啟動 CLI
        try:
            proc = subprocess.Popen(
                cmd,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                bufsize=1,
            )
            yield f"event: log\ndata: {json.dumps({'line':'whisper 已啟動（%s）' % ' '.join(cmd)})}\n\n"
        except Exception as e:
            yield f"event: meta\ndata: {json.dumps({'error': f'whisper 啟動失敗: {e}'})}\n\n"
            yield "event: done\ndata: {}\n\n"
            return

        # 一邊讀取 CLI 輸出（避免 buffer 塞滿），一邊 tail SRT
        index_re = re.compile(r"^\d+$")
        time_re = re.compile(r"\d{2}:\d{2}:\d{2},\d{3}\s+-->\s+\d{2}:\d{2}:\d{2},\d{3}")
        buffer: list[str] = []
        srt_f = None
        last_heartbeat = time.time()

        def flush_buffer():
            nonlocal buffer
            if len(buffer) >= 2:
                tmp = buffer[:]
                if index_re.match(tmp[0].strip()) and len(tmp) >= 3 and time_re.match(tmp[1]):
                    tmp = tmp[1:]
                if time_re.match(tmp[0]):
                    ts = tmp[0]
                    text = " ".join(tmp[1:]).strip()
                    payload = {"ts": ts, "text": text}
                    yield f"event: segment\ndata: {json.dumps(payload, ensure_ascii=False)}\n\n"
            buffer = []

        try:
            while True:
                # keep-alive
                now = time.time()
                if now - last_heartbeat >= 10:
                    yield ": heartbeat\n\n"
                    last_heartbeat = now

                # 開 SRT 檔
                if srt_f is None and os.path.exists(srt_path):
                    srt_f = open(srt_path, "r", encoding="utf-8", errors="ignore")

                progressed = False

                # 讀 SRT
                if srt_f is not None:
                    line = srt_f.readline()
                    if line:
                        progressed = True
                        line = line.rstrip("\n")
                        if not line.strip():
                            for out in flush_buffer():
                                yield out
                        else:
                            buffer.append(line)

                # 吸走 CLI stdout（同時解析即時段落）
                if proc.stdout and not proc.stdout.closed:
                    try:
                        cli_line = proc.stdout.readline()
                        if cli_line:
                            progressed = True
                            ln = cli_line.rstrip("\n")
                            # 解析像：[00:00:00.000 --> 00:00:04.000]  文字
                            m = re.match(r"^\[(\d{2}:\d{2}:\d{2}\.\d{3})\s*-->\s*(\d{2}:\d{2}:\d{2}\.\d{3})\]\s*(.+)$", ln)
                            if m:
                                ts = f"{m.group(1).replace('.', ',')} --> {m.group(2).replace('.', ',')}"
                                text = m.group(3).strip()
                                yield f"event: segment\ndata: {json.dumps({'ts': ts, 'text': text}, ensure_ascii=False)}\n\n"
                            # 若你想在前端顯示所有日誌，可打開下行
                            # else:
                            #     yield f"event: log\ndata: {json.dumps({'line': ln})}\n\n"
                    except Exception:
                        pass

                # 進程是否結束
                if proc.poll() is not None:
                    # 刷尾端
                    if buffer:
                        for out in flush_buffer():
                            yield out
                    break

                if not progressed:
                    time.sleep(0.15)

        finally:
            try:
                if srt_f:
                    srt_f.close()
            except Exception:
                pass

        # ===== 產出 polished 檔案與下載連結 =====
        try:
            if not os.path.exists(txt_raw_path):
                raise FileNotFoundError(f"raw txt not found: {txt_raw_path}")

            with open(txt_raw_path, "r", encoding="utf-8", errors="ignore") as f:
                raw_text = f.read().strip()

            cleaned_text = clean_text(raw_text)
            traditional_text = _converter.convert(cleaned_text)

            polished_text = ""
            try:
                polished_text = punctuate_with_gpt(traditional_text)
            except Exception:
                try:
                    polished_text = bart_punctuation(traditional_text)
                except Exception:
                    polished_text = basic_punctuation(traditional_text)

            if len(polished_text.strip()) < 50 and len(traditional_text.strip()) >= 50:
                polished_text = basic_punctuation(traditional_text)

            try:
                srt_single_lines = extract_srt_lines_with_timestamp(file_path)
            except Exception:
                srt_single_lines = []

            lines_punct = []
            if srt_single_lines:
                try:
                    lines_punct = punctuate_srt_lines(srt_single_lines, batch_size=30)
                except Exception:
                    lines_punct = []

            if len(polished_text.strip()) < 20 and (lines_punct or srt_single_lines):
                merged = _merge_lines_text(lines_punct or srt_single_lines)
                if len(merged) > len(polished_text):
                    polished_text = merged

            target_dir = _meeting_dir(meeting_id)
            base_name = os.path.basename(base)

            outputs = []

            # txt
            txt_filename = f"{base_name}_polished.txt"
            with open(os.path.join(target_dir, txt_filename), "w", encoding="utf-8") as f:
                f.write(polished_text)
            outputs.append(("逐字稿", txt_filename, f"meeting_{meeting_id}/{txt_filename}"))

            # docx
            docx_filename = f"{base_name}_polished.docx"
            docx_path = os.path.join(target_dir, docx_filename)
            doc = Document()
            doc.add_heading("逐字稿紀錄", level=1)
            for para in polished_text.splitlines():
                doc.add_paragraph(para)
            doc.save(docx_path)
            outputs.append(("逐字稿", docx_filename, f"meeting_{meeting_id}/{docx_filename}"))

            # 單行 srt
            single_lines_to_write = lines_punct if lines_punct else srt_single_lines
            if single_lines_to_write:
                single_filename = f"{base_name}_srt_single.txt"
                with open(os.path.join(target_dir, single_filename), "w", encoding="utf-8") as f:
                    f.write("\n".join(single_lines_to_write).rstrip() + "\n")
                outputs.append(("字幕單行", single_filename, f"meeting_{meeting_id}/{single_filename}"))

                # 標準 srt
                try:
                    srt_text = list_to_srt_blocks(single_lines_to_write)
                except Exception:
                    srt_text = write_entries(parse_lines("\n".join(single_lines_to_write) + "\n"), keep_custom=False)
                srt_filename = f"{base_name}_standard.srt"
                with open(os.path.join(target_dir, srt_filename), "w", encoding="utf-8") as f:
                    f.write(srt_text)
                outputs.append(("字幕SRT", srt_filename, f"meeting_{meeting_id}/{srt_filename}"))

            # DB upsert
            for file_type, fname, rel in outputs:
                _upsert_file_record(meeting_id, fname, rel, file_type, uploaded_by=uploader_id)

            # 回傳下載連結
            meta = {"download_txt": f"/download/{meeting_id}/{txt_filename}"}
            yield f"event: meta\ndata: {json.dumps(meta, ensure_ascii=False)}\n\n"

        except Exception as e:
            yield f"event: meta\ndata: {json.dumps({'error': str(e)})}\n\n"

        yield "event: done\ndata: {}\n\n"

    headers = {
        "Content-Type": "text/event-stream; charset=utf-8",
        "Cache-Control": "no-cache, no-transform",
        "Connection": "keep-alive",
        "X-Accel-Buffering": "no",  # 若有 Nginx，請確認也有這個 header
    }
    return Response(stream_with_context(gen()), headers=headers)

# ========== Whisper 轉文字（一次性，非串流） ==========
@speech_bp.route("/whisper", methods=["POST"])
def whisper_transcribe():
    """
    備用的一次性 API（非串流）。
    """
    try:
        data = request.get_json(silent=True) or {}
        filename   = data.get("filename")
        use_gpt    = bool(data.get("use_gpt", True))
        use_basic  = bool(data.get("use_basic", False))
        srt_punct  = bool(data.get("srt_punct", True))
        srt_stdout = bool(data.get("srt_standard", True))

        meeting_id = data.get("meeting_id")
        if not meeting_id:
            return jsonify({"success": False, "error": "未提供 meeting_id"}), 400
        try:
            meeting_id = int(meeting_id)
        except Exception:
            return jsonify({"success": False, "error": "meeting_id 必須是整數"}), 400

        if not filename:
            return jsonify({"success": False, "error": "未提供檔名"}), 400

        user_info = session.get("user") or {}
        uploader_id = user_info.get("id")
        if uploader_id is None:
            return jsonify({"success": False, "error": "未登入，無法建立檔案紀錄"}), 401
        try:
            uploader_id = int(uploader_id)
        except Exception:
            uploader_id = 0

        file_path = os.path.join(UPLOAD_FOLDER, filename)
        if not os.path.exists(file_path):
            return jsonify({"success": False, "error": "檔案不存在"}), 404

        # 1) 轉錄
        transcribe_audio_with_cli(file_path, model_size="medium", lang="zh", speed_profile="balanced")

        # 2) 讀 raw txt
        txt_raw_path = os.path.splitext(file_path)[0] + ".txt"
        if not os.path.exists(txt_raw_path):
            return jsonify({"success": False, "error": f"找不到 Whisper CLI 輸出的 txt 檔案：{txt_raw_path}"}), 500
        with open(txt_raw_path, "r", encoding="utf-8") as f:
            raw_text = f.read().strip()

        cleaned_text = clean_text(raw_text)
        traditional_text = _converter.convert(cleaned_text)

        # 3) 補標點
        polished_text = ""
        if use_gpt:
            try:
                polished_text = punctuate_with_gpt(traditional_text)
            except Exception:
                pass
        if not polished_text:
            if use_basic:
                polished_text = basic_punctuation(traditional_text)
            else:
                try:
                    polished_text = bart_punctuation(traditional_text)
                except Exception:
                    polished_text = basic_punctuation(traditional_text)
        if len(polished_text.strip()) < 50 and len(traditional_text.strip()) >= 50:
            polished_text = basic_punctuation(traditional_text)

        # 4) 單行 SRT
        try:
            srt_single_lines = extract_srt_lines_with_timestamp(file_path)
        except Exception:
            srt_single_lines = []

        lines_punct = []
        if srt_punct and srt_single_lines:
            try:
                lines_punct = punctuate_srt_lines(srt_single_lines, batch_size=30)
            except Exception:
                lines_punct = []

        if len(polished_text.strip()) < 20 and srt_single_lines:
            merged = _merge_lines_text(lines_punct or srt_single_lines)
            if len(merged) > len(polished_text):
                polished_text = merged

        # 5) 寫檔
        target_dir = _meeting_dir(meeting_id)
        base_name = os.path.splitext(filename)[0]

        outputs = []

        txt_filename = f"{base_name}_polished.txt"
        with open(os.path.join(target_dir, txt_filename), "w", encoding="utf-8") as f:
            f.write(polished_text)
        outputs.append(("逐字稿", txt_filename, f"meeting_{meeting_id}/{txt_filename}"))

        docx_filename = f"{base_name}_polished.docx"
        docx_path = os.path.join(target_dir, docx_filename)
        doc = Document()
        doc.add_heading("逐字稿紀錄", level=1)
        for para in polished_text.splitlines():
            doc.add_paragraph(para)
        doc.save(docx_path)
        outputs.append(("逐字稿", docx_filename, f"meeting_{meeting_id}/{docx_filename}"))

        single_filename = f"{base_name}_srt_single.txt"
        single_lines_to_write = lines_punct if lines_punct else srt_single_lines
        if single_lines_to_write:
            with open(os.path.join(target_dir, single_filename), "w", encoding="utf-8") as f:
                f.write("\n".join(single_lines_to_write).rstrip() + "\n")
            outputs.append(("字幕單行", single_filename, f"meeting_{meeting_id}/{single_filename}"))

        srt_filename = None
        if srt_stdout and single_lines_to_write:
            try:
                srt_text = list_to_srt_blocks(single_lines_to_write)
            except Exception:
                srt_text = write_entries(parse_lines("\n".join(single_lines_to_write) + "\n"), keep_custom=False)
            srt_filename = f"{base_name}_standard.srt"
            with open(os.path.join(target_dir, srt_filename), "w", encoding="utf-8") as f:
                f.write(srt_text)
            outputs.append(("字幕SRT", srt_filename, f"meeting_{meeting_id}/{srt_filename}"))

        for file_type, fname, rel in outputs:
            _upsert_file_record(meeting_id, fname, rel, file_type, uploaded_by=uploader_id)

        resp = {
            "success": True,
            "result": polished_text,
            "lines": srt_single_lines,
            "lines_punct": lines_punct,
            "download_txt": f"/download/{meeting_id}/{txt_filename}",
            "download_docx": f"/download/{meeting_id}/{docx_filename}",
            "len_result": len(polished_text),
            "len_raw": len(traditional_text),
            "meeting_folder": f"uploads/meeting_{meeting_id}",
            "files_meta": [
                {"file_name": fname, "file_path": rel, "file_type": ftype}
                for (ftype, fname, rel) in outputs
            ],
        }
        if single_lines_to_write:
            resp["download_srt_single"] = f"/download/{meeting_id}/{single_filename}"
        if srt_filename:
            resp["download_srt_standard"] = f"/download/{meeting_id}/{srt_filename}"

        return jsonify(resp), 200

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500

# ========== 儲存使用者編輯後的段落（產生 TXT / SRT / DOCX + DB 紀錄）==========
@speech_bp.route("/transcript/save", methods=["POST"])
def save_edited_transcript():
    """
    前端送來：
    {
      "meeting_id": 123,
      "src_filename": "20240820_foo.wav",     // 可選，用於命名
      "segments": [{"ts":"00:00:05,000 --> 00:00:06,000","speaker":"說話者1","text":"感謝觀看"}, ...]
    }
    產出：
      - *_with_speakers.txt / .docx / .srt（若 ts 具 '-->'）
    """
    try:
        payload = request.get_json(silent=True) or {}
        meeting_id = payload.get("meeting_id")
        segments = payload.get("segments") or []
        src_filename = payload.get("src_filename") or "edited"

        if not meeting_id:
            return jsonify({"success": False, "error": "未提供 meeting_id"}), 400
        try:
            meeting_id = int(meeting_id)
        except Exception:
            return jsonify({"success": False, "error": "meeting_id 必須是整數"}), 400

        if not isinstance(segments, list) or not segments:
            return jsonify({"success": False, "error": "segments 必須是非空陣列"}), 400

        user_info = session.get("user") or {}
        uploader_id = user_info.get("id")
        if uploader_id is None:
            return jsonify({"success": False, "error": "未登入，無法寫入檔案"}), 401
        try:
            uploader_id = int(uploader_id)
        except Exception:
            uploader_id = 0

        target_dir = _meeting_dir(meeting_id)

        base = os.path.splitext(src_filename)[0] or f"meeting_{meeting_id}"
        prefix = f"{base}_with_speakers"

        # 1) TXT
        txt_lines: list[str] = []
        has_ts = False
        for seg in segments:
            ts = (seg.get("ts") or "").strip()
            sp = (seg.get("speaker") or "").strip() or "說話者"
            tx = (seg.get("text") or "").strip()
            if ts:
                has_ts = has_ts or ("-->" in ts)
                txt_lines.append(f"{ts} {sp}：{tx}")
            else:
                txt_lines.append(f"{sp}：{tx}")

        txt_name = f"{prefix}.txt"
        txt_path = os.path.join(target_dir, txt_name)
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("\n".join(txt_lines).rstrip() + "\n")
        _upsert_file_record(meeting_id, txt_name, f"meeting_{meeting_id}/{txt_name}", "逐字稿(帶說話者)", uploaded_by=uploader_id)

        # 2) DOCX
        docx_name = f"{prefix}.docx"
        docx_path = os.path.join(target_dir, docx_name)
        doc = Document()
        doc.add_heading("逐字稿（含說話者）", level=1)
        for line in txt_lines:
            doc.add_paragraph(line)
        doc.save(docx_path)
        _upsert_file_record(meeting_id, docx_name, f"meeting_{meeting_id}/{docx_name}", "逐字稿DOCX(帶說話者)", uploaded_by=uploader_id)

        # 3) SRT（有 ts 時）
        srt_name = None
        if has_ts:
            single_lines = []
            for seg in segments:
                ts = (seg.get("ts") or "").strip()
                sp = (seg.get("speaker") or "").strip() or "說話者"
                tx = (seg.get("text") or "").strip()
                if "-->" in ts:
                    single_lines.append(f"{ts} | {sp}：{tx}")
            try:
                srt_text = list_to_srt_blocks(single_lines)
            except Exception:
                srt_text = write_entries(parse_lines("\n".join(single_lines) + "\n"), keep_custom=False)

            srt_name = f"{prefix}.srt"
            srt_path = os.path.join(target_dir, srt_name)
            with open(srt_path, "w", encoding="utf-8") as f:
                f.write(srt_text)
            _upsert_file_record(meeting_id, srt_name, f"meeting_{meeting_id}/{srt_name}", "字幕SRT(帶說話者)", uploaded_by=uploader_id)

        resp = {
            "success": True,
            "download_txt": f"/download/{meeting_id}/{txt_name}",
            "download_docx": f"/download/{meeting_id}/{docx_name}",
        }
        if srt_name:
            resp["download_srt"] = f"/download/{meeting_id}/{srt_name}"
        return jsonify(resp), 200

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500

# ========== 檔案下載（會議資料夾）==========
@speech_bp.route("/download/<int:meeting_id>/<path:filename>")
def download_file(meeting_id, filename):
    directory = _meeting_dir(meeting_id)
    return send_from_directory(directory, filename, as_attachment=True)
